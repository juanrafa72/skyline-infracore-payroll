import sys, openpyxl
from openpyxl.utils import get_column_letter
path = sys.argv[1]
wb = openpyxl.load_workbook(path)
print("### DEFINED NAMES")
try:
    for name, dn in wb.defined_names.items():
        print("  %s = %s" % (name, dn.value))
except AttributeError:
    for dn in wb.defined_names.definedName:
        print("  %s = %s" % (dn.name, dn.value))
print()
print("### TABLES")
for ws in wb.worksheets:
    try:
        for t in ws.tables.values():
            print("  [%s] %s ref=%s" % (ws.title, t.name, t.ref))
            for col in t.tableColumns:
                print("        - %s  %s" % (col.name, (col.calculatedColumnFormula.attr_text if col.calculatedColumnFormula else "")))
    except Exception as e:
        print("  ERR", ws.title, e)
print()
print("### DATA VALIDATIONS")
for ws in wb.worksheets:
    for dv in ws.data_validations.dataValidation:
        print("  [%s] %s sqref=%s formula1=%s" % (ws.title, dv.type, dv.sqref, dv.formula1))
